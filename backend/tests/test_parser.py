"""Text-extraction tests: messy PDF/DOCX layouts must survive extraction,
and the parser prompt must forbid inventing facts.

Generates 7 messy PDF fixtures + 1 DOCX into tests/sample_resumes/ (created on
first run, committed afterwards). Any extra files dropped into that directory
(e.g. your own real resumes) are picked up automatically and only checked for
non-empty extraction, since their content is unknown.
"""
import asyncio
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import fitz  # PyMuPDF
import docx as docx_lib

from app.services.file_parser import extract_text, extract_text_from_pdf, extract_text_from_docx

FIXTURES = Path(__file__).resolve().parent / "sample_resumes"
failures = []


# ---------------------------------------------------------------- fixture builders

def _pdf(builder) -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4 points
    builder(page)
    data = doc.tobytes()
    doc.close()
    return data


def two_column(page):
    left = [("Priya Raman", 40, 60), ("priya.raman@example.com", 40, 78), ("+44 20 7946 0958", 40, 96)]
    for text, x, y in left:
        page.insert_text((x, y), text, fontsize=10)
    right = [
        ("EXPERIENCE", 300, 60),
        ("Software Engineer, Nimbus Labs", 300, 80),
        ("Built data pipelines in Python", 300, 98),
        ("Led team of 6 engineers", 300, 116),
        ("SKILLS", 300, 150),
        ("Python, SQL, Airflow", 300, 168),
    ]
    for text, x, y in right:
        page.insert_text((x, y), text, fontsize=10)


def table_grid(page):
    page.insert_text((40, 40), "Miguel Santos", fontsize=14)
    page.insert_text((40, 56), "miguel.santos@example.org", fontsize=10)
    page.insert_text((40, 70), "+1 415 555 0132", fontsize=10)
    cols = [40, 240, 420]
    rows = [
        ("ROLE", "COMPANY", "DATES"),
        ("Software Engineer", "Acme Corp", "2019 - 2023"),
        ("DevOps Engineer", "Globex", "2023 - Present"),
        ("Intern", "Initech", "2018"),
    ]
    for r, row in enumerate(rows):
        y = 110 + r * 34
        for c, cell in enumerate(row):
            page.insert_text((cols[c], y), cell, fontsize=9, fontname="helv")
        page.draw_line((40, y + 8), (540, y + 8), width=0.5)


def unicode_text(page):
    lines = [
        "Zoë Müller-Ñandú",
        "zoe.muller@example.de",
        "Straße 45, München",
        "“AI/ML Engineer” — e.g. résumé résumé",
        "Earned €1,200,000 in revenue",
        "José's team: Françoise, Bjørn",
    ]
    for i, line in enumerate(lines):
        page.insert_text((60, 70 + i * 22), line, fontsize=11)


def header_footer(page):
    page.insert_text((40, 40), "Dana Kim", fontsize=16)
    page.insert_text((40, 58), "dana.kim@example.net", fontsize=10)
    page.insert_text((40, 100), "PROFESSIONAL EXPERIENCE", fontsize=12)
    page.insert_text((60, 122), "Senior Analyst at UMBRACorp (2021 - Present)", fontsize=10)
    page.insert_text((60, 142), "Automated reporting, saving 12 hours per week", fontsize=10)
    # rotated sidebar text (90 degrees)
    page.insert_text((20, 300), "CONFIDENTIAL DRAFT", fontsize=9, rotate=90)
    page.insert_text((40, 810), "Page 1 of 2", fontsize=8)
    page.insert_text((480, 810), "updated 05/2024", fontsize=8)


def minimal(page):
    page.insert_text((60, 400), "Sam Okafor", fontsize=12)
    page.insert_text((60, 418), "sam.okafor@mail.test", fontsize=11)


def long_lines(page):
    page.insert_text((40, 60), "Lee Chen", fontsize=12)
    page.insert_text((40, 78), "lee.chen@example.io", fontsize=10)
    page.insert_text((40, 100), "https://portfolio.example.com/very/long/path/segment/one/two/three/four?query=withvalue", fontsize=9)
    page.insert_text((40, 130), "Expert in micro-", fontsize=10)
    page.insert_text((40, 146), "services and distributed systems", fontsize=10)
    page.insert_text((40, 176), "5000000000000000000000 requests", fontsize=10)


def scattered(page):
    page.insert_text((100, 90), "Ava Thompson", fontsize=13)
    page.insert_text((380, 310), "ava.thompson@example.com", fontsize=7)
    page.insert_text((70, 520), "Marketing Lead", fontsize=11)
    page.insert_text((300, 700), "+61 2 5550 8899", fontsize=8)
    page.insert_text((60, 780), "Grew organic traffic 3x", fontsize=10)


PDF_FIXTURES = [
    ("two_column.pdf", two_column, ["Priya Raman", "priya.raman@example.com", "+44 20 7946 0958"]),
    ("table_grid.pdf", table_grid, ["Miguel Santos", "miguel.santos@example.org", "Software Engineer", "+1 415 555 0132"]),
    ("unicode.pdf", unicode_text, ["Zoë Müller-Ñandú", "zoe.muller@example.de", "1,200,000"]),
    ("header_footer.pdf", header_footer, ["Dana Kim", "dana.kim@example.net", "PROFESSIONAL EXPERIENCE", "Page 1 of 2"]),
    ("minimal.pdf", minimal, ["Sam Okafor", "sam.okafor@mail.test"]),
    ("long_lines.pdf", long_lines, ["micro-", "lee.chen@example.io", "5000000000000000000000"]),
    ("scattered.pdf", scattered, ["Ava Thompson", "ava.thompson@example.com", "3x"]),
]

DOCX_NAME = "tables_uni.docx"
DOCX_MARKERS = ["Alice Ferreira", "alice.ferreira@example.com", "+55 11 91234 5678", "APRESENTAÇÃO", "Olá"]


def build_docx() -> bytes:
    d = docx_lib.Document()
    d.add_paragraph("Alice Ferreira")
    d.add_heading("ALICE FERREIRA", level=1)
    d.add_paragraph("alice.ferreira@example.com")
    d.add_paragraph("+55 11 91234 5678")
    d.add_heading("APRESENTAÇÃO", level=2)
    d.add_paragraph("Olá! Engenheira de software com 7 anos de experiência.")
    table = d.add_table(rows=2, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Cargo", "Empresa", "Período"
    row = table.rows[1].cells
    row[0].text, row[1].text, row[2].text = "Engenheira de Software", "Solutions SA", "2020 - Presente"
    d.add_paragraph("")  # empty paragraph soup
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- run fixtures

FIXTURES.mkdir(exist_ok=True)

for name, builder, markers in PDF_FIXTURES:
    path = FIXTURES / name
    if not path.exists():
        path.write_bytes(_pdf(builder))
    text = extract_text_from_pdf(path.read_bytes())
    if not text.strip():
        failures.append(f"{name}: extracted text is empty")
        continue
    for marker in markers:
        if marker not in text:
            failures.append(f"{name}: missing {marker!r} after extraction")
    print(f"{name}: {len(text)} chars, {len(markers)} markers checked")

docx_path = FIXTURES / DOCX_NAME
if not docx_path.exists():
    docx_path.write_bytes(build_docx())
docx_text = extract_text_from_docx(docx_path.read_bytes())
for marker in DOCX_MARKERS:
    if marker not in docx_text:
        failures.append(f"{DOCX_NAME}: missing {marker!r} after extraction")
print(f"{DOCX_NAME}: {len(docx_text)} chars, {len(DOCX_MARKERS)} markers checked")

# ---------------------------------------------------------------- user-dropped resumes

known = {name for name, _, _ in PDF_FIXTURES} | {DOCX_NAME}
extras = [p for p in FIXTURES.iterdir() if p.is_file() and p.name not in known and not p.name.startswith(".")]
for path in extras:
    try:
        text = extract_text(path.read_bytes(), path.name)
        if not text.strip():
            failures.append(f"{path.name}: extracted text is empty")
        else:
            print(f"{path.name}: {len(text)} chars (user fixture, non-empty OK)")
    except Exception as e:
        failures.append(f"{path.name}: extraction failed: {e}")

# ---------------------------------------------------------------- dispatch + prompt guard

try:
    extract_text(b"not a real file", "resume.txt")
    failures.append("extract_text accepted unsupported .txt file")
except ValueError:
    print("unsupported extension rejected: OK")

# Parser prompt must explicitly forbid inventing facts (regression guard)
from app.agents import parser_agent

_captured = {}
_orig_generate_json = parser_agent.llm.generate_json


async def _fake_generate_json(system_prompt, user_prompt, **kwargs):
    _captured["system"] = system_prompt
    return {"name": "Test Person", "summary": "Test summary", "email": None, "experiences": [], "skills": []}


async def _run_parse():
    return await parser_agent.parse_resume_text("raw resume text")


parser_agent.llm.generate_json = _fake_generate_json
try:
    parsed = asyncio.run(_run_parse())
finally:
    parser_agent.llm.generate_json = _orig_generate_json

sys_prompt = _captured.get("system", "")
if "Never invent" not in sys_prompt:
    failures.append("parser prompt missing 'Never invent' directive")
if "null" not in sys_prompt:
    failures.append("parser prompt missing null-when-unsure directive")
if parsed.name != "Test Person":
    failures.append("parse_resume_text did not return Resume from LLM JSON")
print("parser prompt anti-hallucination directives: OK")

if failures:
    print("\nFAILURES:")
    for f in failures:
        print(" -", f)
    sys.exit(1)
print("\nALL PARSER/EXTRACTION CHECKS PASSED")
