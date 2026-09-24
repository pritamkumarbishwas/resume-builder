import pymupdf
import docx
import io

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text = ""
    with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text += page.get_text() + "\n"
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
    
def extract_text(file_bytes: bytes, filename: str) -> str:
    if filename.lower().endswith('.pdf'):
        return extract_text_from_pdf(file_bytes)
    elif filename.lower().endswith('.docx'):
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError("Unsupported file format. Please upload a PDF or DOCX file.")
