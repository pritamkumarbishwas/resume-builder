from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from app.schemas.resume import Resume
from jinja2 import Environment, FileSystemLoader
from xhtml2pdf import pisa
from docx import Document
from docx.shared import Pt, Inches, RGBColor
import io
import os

router = APIRouter()

template_dir = os.path.join(os.path.dirname(__file__), '..', '..', 'templates')
env = Environment(loader=FileSystemLoader(template_dir))

def sanitize_text(data):
    if isinstance(data, dict):
        return {k: sanitize_text(v) for k, v in data.items()}
    elif isinstance(data, list):
        return [sanitize_text(i) for i in data]
    elif isinstance(data, str):
        # Replace common unicode characters that xhtml2pdf's default fonts can't render
        return (data.replace('–', '-')
                    .replace('—', '-')
                    .replace('•', '')
                    .replace('’', "'")
                    .replace('‘', "'")
                    .replace('“', '"')
                    .replace('”', '"')
                    .replace('\u200b', '')
                    .replace('…', '...'))
    return data

@router.post("/pdf")
async def export_pdf(resume: Resume, template: str = "classic"):
    try:
        # Fallback to classic if template doesn't exist
        template_file = f'{template}.html'
        try:
            html_template = env.get_template(template_file)
        except Exception:
            html_template = env.get_template('classic.html')
            
        sanitized_data = sanitize_text(resume.model_dump())
        html_out = html_template.render(resume=sanitized_data)
        
        result_file = io.BytesIO()
        pisa_status = pisa.CreatePDF(html_out, dest=result_file, encoding='utf-8')
        
        if pisa_status.err:
            raise Exception("PDF generation failed")
            
        return Response(content=result_file.getvalue(), media_type="application/pdf")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


ACCENT = RGBColor(0x7C, 0x3A, 0xED)


def _section_heading(doc, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = ACCENT
    return p


def _meta_line(doc, parts, italic=True, size=9.5):
    text = " | ".join(str(p) for p in parts if p)
    if not text:
        return None
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(2)
    return p


@router.post("/docx")
async def export_docx(resume: Resume):
    """ATS-friendly single-column DOCX export (no tables, no columns)."""
    try:
        data = resume.model_dump()
        doc = Document()

        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)

        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        name_p = doc.add_paragraph()
        name_run = name_p.add_run(data.get("name") or "")
        name_run.bold = True
        name_run.font.size = Pt(22)
        name_p.paragraph_format.space_after = Pt(2)

        _meta_line(
            doc,
            [data.get("email"), data.get("phone"), data.get("linkedin"), data.get("portfolio")],
            italic=False,
            size=9.5,
        )

        if data.get("summary"):
            _section_heading(doc, "Professional Summary")
            p = doc.add_paragraph(data["summary"])
            p.paragraph_format.space_after = Pt(4)

        if data.get("experiences"):
            _section_heading(doc, "Experience")
            for exp in data["experiences"]:
                title_p = doc.add_paragraph()
                title_p.paragraph_format.space_before = Pt(6)
                title_p.paragraph_format.space_after = Pt(0)
                title_run = title_p.add_run(exp.get("title") or "")
                title_run.bold = True
                title_run.font.size = Pt(11)
                _meta_line(
                    doc,
                    [
                        exp.get("company"),
                        exp.get("location"),
                        f"{exp.get('start_date') or ''} - {exp.get('end_date') or 'Present'}".strip(" -"),
                    ],
                )
                for bullet in exp.get("description") or []:
                    if not bullet.strip():
                        continue
                    bp = doc.add_paragraph(bullet, style="List Bullet")
                    bp.paragraph_format.space_after = Pt(0)
                    bp.paragraph_format.left_indent = Inches(0.25)

        if data.get("projects"):
            _section_heading(doc, "Projects")
            for project in data["projects"]:
                if not (project.get("name") or project.get("description")):
                    continue
                proj_p = doc.add_paragraph()
                proj_p.paragraph_format.space_before = Pt(6)
                proj_p.paragraph_format.space_after = Pt(0)
                name_run = proj_p.add_run(project.get("name") or "")
                name_run.bold = True
                name_run.font.size = Pt(11)
                _meta_line(doc, [", ".join(project.get("technologies") or []), project.get("link")])
                for bullet in project.get("description") or []:
                    if not bullet.strip():
                        continue
                    bp = doc.add_paragraph(bullet, style="List Bullet")
                    bp.paragraph_format.space_after = Pt(0)
                    bp.paragraph_format.left_indent = Inches(0.25)

        if data.get("education"):
            _section_heading(doc, "Education")
            for edu in data["education"]:
                edu_p = doc.add_paragraph()
                edu_p.paragraph_format.space_before = Pt(4)
                edu_p.paragraph_format.space_after = Pt(0)
                deg_run = edu_p.add_run(edu.get("degree") or "")
                deg_run.bold = True
                deg_run.font.size = Pt(11)
                _meta_line(
                    doc,
                    [edu.get("institution"), edu.get("location"), edu.get("graduation_date"), edu.get("gpa") and f"GPA: {edu['gpa']}"],
                )

        if data.get("skills"):
            _section_heading(doc, "Skills")
            for group in data["skills"]:
                skills = [s for s in (group.get("skills") or []) if s]
                if not skills and not group.get("category"):
                    continue
                gp = doc.add_paragraph()
                gp.paragraph_format.space_after = Pt(2)
                cat_run = gp.add_run(f"{group.get('category') or 'Skills'}: ")
                cat_run.bold = True
                gp.add_run(", ".join(skills))

        buffer = io.BytesIO()
        doc.save(buffer)
        return Response(
            content=buffer.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="Tailored_Resume.docx"'},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
